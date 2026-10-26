"""PDF text extraction and statute-aware chunking with page and section provenance."""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

import fitz  # PyMuPDF

from core.errors import InputValidationError
from core.models import Chunk
from retrieval.legal_refs import section_key

# Phrases that look like instructions to an AI system. Retrieved text is evidence, never instructions;
# chunks containing these patterns are flagged and neutralised before being shown to the LLM.
INJECTION_PATTERNS = [
    r"ignore\s+(all\s+|any\s+)?(previous|prior|above)\s+instructions",
    r"disregard\s+(the\s+)?(system|previous)\s+(prompt|instructions)",
    r"you\s+are\s+now\s+(a|an|in)\b",
    r"\bsystem\s*prompt\b",
    r"<\|?(im_start|im_end|system|assistant)\|?>",
    r"\bact\s+as\s+(an?\s+)?(unrestricted|jailbroken)",
    r"(mark|label|classify)\s+(this|the)\s+(answer|claim)s?\s+as\s+(verified|supported|safe)",
]
_INJECTION_RE = re.compile("|".join(INJECTION_PATTERNS), re.I)

# Amendment footnotes in the Constitution: "1. Subs. by the Constitution ..."
_FOOTNOTE_RE = re.compile(
    r"^\s*\d{1,3}\.?\s+(Subs|Ins|Omitted|Added|Rep|The words|Cl\.|Art\.|Renumbered|See|Sub-s|Now|Re-numbered)\b", re.I)
# Section/article starts: "103. (1) Whoever", "2[21A. Right to", "66C  ", "70 A  National", "66E.  "
_START_RE = re.compile(
    r"^\s*(?:\d{0,2}\[)?(?P<num>\d{1,3}(?:\s?[A-Z]{1,3})?)(?:\.(?!\d)\s*(?P<rest1>.*)|(?P<rest2>\s+[A-Z(“\"].*|\s*$))")


@dataclass
class PageText:
    number: int   # 1-based
    text: str


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


SUPPORTED_EXTS = {".pdf", ".txt", ".docx"}


def detect_file_type(path: Path) -> str:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise InputValidationError(f"Unsupported file type '{ext}'. Accepted: PDF, TXT, DOCX.")
    return ext.lstrip(".")


def validate_pdf(path: Path, max_mb: int, max_pages: int) -> int:
    """Validate an uploaded PDF/TXT/DOCX: type, magic bytes, size, page count. Returns page/unit count."""
    ext = detect_file_type(path)
    size_mb = path.stat().st_size / (1 << 20)
    if size_mb > max_mb:
        raise InputValidationError(f"File is {size_mb:.1f} MB; the limit is {max_mb} MB.")
    with open(path, "rb") as fh:
        head = fh.read(8)
    if ext == "pdf":
        if head[:5] != b"%PDF-":
            raise InputValidationError("File does not have a valid PDF header.")
        try:
            with fitz.open(path) as doc:
                if doc.needs_pass:
                    raise InputValidationError("Encrypted PDFs are not supported.")
                pages = doc.page_count
        except InputValidationError:
            raise
        except Exception as exc:
            raise InputValidationError(f"PDF could not be parsed: {exc}") from exc
        if pages == 0 or pages > max_pages:
            raise InputValidationError(f"PDF has {pages} pages; the allowed range is 1-{max_pages}.")
        return pages
    if ext == "docx":
        if head[:2] != b"PK":  # docx is a zip container
            raise InputValidationError("File does not look like a valid .docx (zip) document.")
        return 1
    return 1  # txt


def _text_density(doc) -> float:
    """Average extractable characters per page; low values indicate a scanned/image-only PDF."""
    if doc.page_count == 0:
        return 0.0
    total = sum(len(doc[i].get_text("text").strip()) for i in range(min(doc.page_count, 20)))
    return total / min(doc.page_count, 20)


def assert_extractable_pdf(path: Path, min_chars_per_page: int = 80) -> None:
    """Raise a clear error for scanned/image-only PDFs rather than indexing empty text."""
    with fitz.open(path) as doc:
        density = _text_density(doc)
    if density < min_chars_per_page:
        raise InputValidationError(
            f"This PDF appears to be scanned or image-only (~{density:.0f} extractable chars/page). "
            "OCR it before uploading; TrustGuard does not index image-only documents.")


def _is_latin_line(line: str) -> bool:
    letters = [c for c in line if c.isalpha()]
    if not letters:
        return True
    return sum(c.isascii() for c in letters) / len(letters) > 0.8


_LEGACY_FONT_RE = re.compile(r"[À-ÿŒœˆ˜‘-„¡-¿]{3,}")


def _clean_line(line: str) -> str:
    line = line.replace(" ", " ").rstrip()
    # drop lines that are mostly legacy-font (Kruti Dev) Hindi rendered as Latin-1 garbage
    if _LEGACY_FONT_RE.search(line) and sum(ch in "¼½¾ÉÊËÌÍ" for ch in line) >= 2:
        return ""
    return line


MARGIN_TAG = "\x00MARGIN:"
_MARGIN_NOISE_RE = re.compile(r"^(\d+\s+of\s+\d{4}\.?|Sec\.\s*\d+\]?)$", re.I)


def _block_lines(text: str) -> list[str]:
    lines = [_clean_line(l) for l in text.splitlines()]
    return [l for l in lines if l.strip() and _is_latin_line(l)]


def extract_text_pages(path: Path) -> list["PageText"]:
    """Extract pages from a TXT or DOCX file as a single logical page (no gazette layout to parse)."""
    ext = path.suffix.lower()
    if ext == ".txt":
        raw = path.read_text(encoding="utf-8", errors="replace")
    elif ext == ".docx":
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover
            raise InputValidationError("python-docx is not installed; cannot read .docx files.") from exc
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:  # preserve table text row by row
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        raw = "\n".join(parts)
    else:
        raise InputValidationError(f"extract_text_pages does not handle '{ext}'")
    lines = [l for l in (_clean_line(x) for x in raw.splitlines()) if l.strip() and _is_latin_line(l)]
    if not lines:
        raise InputValidationError("No extractable text found in the document.")
    return [PageText(number=1, text="\n".join(lines))]


def extract_any(path: Path) -> list["PageText"]:
    """Dispatch to the right extractor by file type; detect scanned PDFs up front."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        assert_extractable_pdf(path)
        return extract_pages(path)
    return extract_text_pages(path)


def extract_pages(path: Path) -> list[PageText]:
    """Extract page text in reading order.

    Gazette statutes print section headings as narrow marginal notes ("Punishment for murder.").
    These blocks are detected geometrically and emitted as tagged lines immediately before the body
    block they sit beside, so the chunker can attach them to the right section as its heading.
    """
    pages = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc):
            width = page.rect.width
            body, margin = [], []
            for b in page.get_text("blocks"):
                x0, y0, x1, y1, text = b[:5]
                if b[6] != 0 or not text.strip():  # skip image blocks
                    continue
                narrow = (x1 - x0) < 0.2 * width
                if narrow and (x1 < 0.22 * width or x0 > 0.78 * width):
                    note = " ".join(_block_lines(text)).strip()
                    if note and not _MARGIN_NOISE_RE.match(note):
                        margin.append((y0, note))
                else:
                    body.append((y0, x0, text))
            body.sort(key=lambda t: (round(t[0]), t[1]))
            margin.sort()
            lines: list[str] = []
            for y0, _, text in body:
                while margin and margin[0][0] <= y0 + 8:
                    lines.append(MARGIN_TAG + margin.pop(0)[1])
                lines.extend(_block_lines(text))
            lines.extend(MARGIN_TAG + m[1] for m in margin)
            pages.append(PageText(number=i + 1, text="\n".join(lines)))
    return pages


@dataclass
class _Start:
    pos: int
    page: int
    num: str
    key: tuple[int, str]
    first_line: str
    margin: str = ""


def _candidate_starts(pages: list[PageText]) -> tuple[str, list[int], list[_Start]]:
    full, page_of_offset, starts = [], [], []
    offset = 0
    # Printed page numbers sit at a page edge and differ from the PDF page index by a constant offset.
    offsets: dict[int, int] = {}
    for p in pages:
        page_lines = p.text.split("\n")
        for idx, line in enumerate(page_lines):
            m = re.match(r"^\s*(\d{1,3})\b", line)
            if m and (idx <= 2 or idx >= len(page_lines) - 2):
                d = int(m.group(1)) - p.number
                offsets[d] = offsets.get(d, 0) + 1
    page_offset = max(offsets, key=offsets.get) if offsets else None
    pending_margin = ""
    for p in pages:
        page_lines = p.text.split("\n")
        for idx, line in enumerate(page_lines):
            if line.startswith(MARGIN_TAG):  # marginal heading: attach to the next section start
                pending_margin = line[len(MARGIN_TAG):]
                continue
            m = _START_RE.match(line)
            if m and not _FOOTNOTE_RE.match(line):
                num = re.sub(r"\s+", "", m.group("num")).upper()
                rest = (m.group("rest1") or m.group("rest2") or "").strip()
                # marginal notes may precede the printed page number, so position is not required
                is_page_number = num.isdigit() and int(num) - p.number == page_offset \
                    and (not rest or re.match(r"(THE GAZETTE|_{3,}|\[Part)", rest))
                if not is_page_number and not re.match(r"(THE GAZETTE|_{3,})", rest):
                    starts.append(_Start(offset, p.number, num, section_key(num), rest, pending_margin))
                    pending_margin = ""
            full.append(line)
            page_of_offset.append(p.number)
            offset += len(line) + 1
    return "\n".join(full), page_of_offset, starts


def _best_run(starts: list[_Start], text_len: int, window: int = 400) -> list[_Start]:
    """Find the numbering chain that best explains the body of the statute.

    Dynamic programme over candidate starts: a chain must strictly increase in section number;
    each link scores the text it covers (capped, so table-of-contents entries score little) and is
    penalised for skipped numbers (so stray footnote numbers cannot hijack the chain).
    """
    n = len(starts)
    if n == 0:
        return []
    neg = float("-inf")
    best = [neg] * n
    parent = [-1] * n
    for i, s in enumerate(starts):
        # any section may begin a chain (uploaded fragments need not start at section 1); a chain that
        # starts at a low number is preferred via a mild seed penalty proportional to the start number.
        best[i] = -60.0 * min(s.key[0], 12)
        for j in range(max(0, i - window), i):
            if best[j] == neg:
                continue
            p = starts[j]
            gap_num = s.key[0] - p.key[0]
            if not (p.key < s.key) or gap_num > 8:
                continue
            # +150 per link rewards including every genuine section rather than jumping over them
            score = best[j] + min(s.pos - p.pos, 4000) + 150 - 400 * max(0, gap_num - 1)
            if score > best[i]:
                best[i], parent[i] = score, j
    end = max(range(n), key=lambda k: best[k])
    chain = []
    while end != -1:
        chain.append(starts[end])
        end = parent[end]
    chain.reverse()
    return chain if len(chain) >= 3 else []


def _heading(first_line: str, body: str) -> str:
    if not first_line:  # heading printed on the line after the number (IT Act layout)
        nxt = body.split("\n")[1:2]
        first_line = nxt[0].strip() if nxt else ""
    first_line = re.sub(r"\(\s*(Inserted|Substituted|Introduced|Amended)\s+[Vv]ide\s+ITA+\s*2008\s*\)", "", first_line).strip()
    m = re.match(r"(.{3,140}?)\.?\s*[—–-]{1,2}", first_line)
    if m and not first_line.lstrip().startswith("("):
        return m.group(1).strip(" .")
    if first_line and not first_line.startswith("(") and len(first_line) < 120 and not first_line.lower().startswith("whoever"):
        return first_line.strip(" .")
    return ""


def chunk_document(doc_id: str, pages: list[PageText], act_code: str, ref_type: str,
                   chunk_chars: int = 1200, overlap: int = 150) -> list[Chunk]:
    full, page_of_line, starts = _candidate_starts(pages)
    line_offsets, off = [], 0
    for line in full.split("\n"):
        line_offsets.append(off)
        off += len(line) + 1

    def page_at(pos: int) -> int:
        lo, hi = 0, len(line_offsets) - 1
        while lo < hi:
            mid = (lo + hi + 1) // 2
            if line_offsets[mid] <= pos:
                lo = mid
            else:
                hi = mid - 1
        return page_of_line[lo] if page_of_line else 1

    run = _best_run(starts, len(full))
    segments: list[tuple[int, int, str, str]] = []  # (start, end, section, heading)
    if run:
        if run[0].pos > 0:
            segments.append((0, run[0].pos, "", "Front matter"))
        for i, s in enumerate(run):
            end = run[i + 1].pos if i + 1 < len(run) else len(full)
            heading = s.margin.strip(" .") or _heading(s.first_line, full[s.pos:end])
            segments.append((s.pos, end, s.num, heading))
    else:
        segments.append((0, len(full), "", ""))

    chunks: list[Chunk] = []
    for seg_start, seg_end, section, heading in segments:
        text = full[seg_start:seg_end]
        if not text.strip():
            continue
        pieces, pos = [], 0
        while pos < len(text):
            end = min(len(text), pos + chunk_chars)
            if end < len(text):  # break on a sentence or line boundary
                cut = max(text.rfind(". ", pos + chunk_chars // 2, end), text.rfind("\n", pos + chunk_chars // 2, end))
                if cut > pos:
                    end = cut + 1
            pieces.append((pos, end))
            if end >= len(text):
                break
            pos = max(end - overlap, pos + 1)
        for a, b in pieces:
            body = re.sub(r"[ \t]+", " ", text[a:b]).strip()
            if len(body) < 25:
                continue
            ordinal = len(chunks)
            chunks.append(Chunk(
                chunk_id=f"{doc_id}:{ordinal:05d}",
                doc_id=doc_id,
                text=body,
                page_start=page_at(seg_start + a),
                page_end=page_at(seg_start + b - 1),
                section=section,
                ref_type=ref_type,
                heading=heading,
                act_code=act_code,
                ordinal=ordinal,
                injection_flag=bool(_INJECTION_RE.search(body)),
            ))
    return chunks


def neutralise_injection(text: str) -> str:
    return _INJECTION_RE.sub("[instruction-like text removed]", text)
